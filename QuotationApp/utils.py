from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_excel_bill(items, gst_rate=18):
    """
    Generate Excel bill with items, quantities, prices, and GST calculations.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    
    # Set column widths
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    
    # Define styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    total_font = Font(bold=True, size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Title
    ws['A1'] = "QUOTATION"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:F1')
    ws['A1'].alignment = Alignment(horizontal='center')
    
    # Headers
    headers = ['Item Name', 'Quantity', 'Unit Price', 'Discount %', 'Total Price', 'GST (%)']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
    
    # Data rows
    row = 4
    subtotal = 0
    
    for item in items:
        quantity = item.get('quantity', 1)
        unit_price = float(item.get('price', 0))
        discount = item.get('discount', 0)
        
        # Calculate price after discount
        discounted_price = unit_price * (1 - discount / 100)
        total_price = quantity * discounted_price
        subtotal += total_price
        
        # Item Name
        ws.cell(row=row, column=1).value = item.get('name', '')
        # Quantity
        ws.cell(row=row, column=2).value = quantity
        # Unit Price
        ws.cell(row=row, column=3).value = unit_price
        ws.cell(row=row, column=3).number_format = '₹#,##0.00'
        # Discount %
        ws.cell(row=row, column=4).value = discount
        # Total Price (after discount)
        ws.cell(row=row, column=5).value = total_price
        ws.cell(row=row, column=5).number_format = '₹#,##0.00'
        # GST %
        ws.cell(row=row, column=6).value = gst_rate
        
        # Apply borders and alignment
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            if col in [2, 3, 4, 5, 6]:
                cell.alignment = Alignment(horizontal='right')
        
        row += 1
    
    # Add blank row
    row += 1
    
    # Subtotal
    ws.cell(row=row, column=4).value = "Subtotal:"
    ws.cell(row=row, column=4).font = total_font
    ws.cell(row=row, column=4).alignment = Alignment(horizontal='right')
    ws.cell(row=row, column=5).value = subtotal
    ws.cell(row=row, column=5).number_format = '₹#,##0.00'
    ws.cell(row=row, column=5).font = total_font
    ws.cell(row=row, column=5).fill = total_fill
    ws.cell(row=row, column=5).border = border
    
    row += 1
    
    # GST
    gst_amount = (subtotal * gst_rate) / 100
    ws.cell(row=row, column=4).value = f"GST ({gst_rate}%):"
    ws.cell(row=row, column=4).font = total_font
    ws.cell(row=row, column=4).alignment = Alignment(horizontal='right')
    ws.cell(row=row, column=5).value = gst_amount
    ws.cell(row=row, column=5).number_format = '₹#,##0.00'
    ws.cell(row=row, column=5).font = total_font
    ws.cell(row=row, column=5).fill = total_fill
    ws.cell(row=row, column=5).border = border
    
    row += 1
    
    # Total
    total = subtotal + gst_amount
    ws.cell(row=row, column=4).value = "TOTAL:"
    ws.cell(row=row, column=4).font = Font(bold=True, size=12, color="FFFFFF")
    ws.cell(row=row, column=4).fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws.cell(row=row, column=4).alignment = Alignment(horizontal='right')
    ws.cell(row=row, column=5).value = total
    ws.cell(row=row, column=5).number_format = '₹#,##0.00'
    ws.cell(row=row, column=5).font = Font(bold=True, size=12, color="FFFFFF")
    ws.cell(row=row, column=5).fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws.cell(row=row, column=5).border = border
    
    # Save file
    filename = 'quotation.xlsx'
    wb.save(filename)
    return filename
